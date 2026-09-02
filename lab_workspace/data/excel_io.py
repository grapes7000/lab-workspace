from pathlib import Path

def export_workbook(db,path):
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill
 wb=Workbook();wb.remove(wb.active)
 for table,title in (("materials","Materials"),("samples","Samples")):
  ws=wb.create_sheet(title); rows=db.search(table,'')
  if rows:
   headers=list(rows[0].keys());ws.append(headers)
   for cell in ws[1]:cell.font=Font(bold=True);cell.fill=PatternFill('solid',fgColor='D9EAF7')
   for row in rows:ws.append([row[h] for h in headers])
   ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
   for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(30,max(12,max(len(str(c.value or '')) for c in col)+2))
 meta=wb.create_sheet('Metadata');meta.append(['Application','Lab Workspace']);meta.append(['Schema','v1.2'])
 wb.save(path)

def import_sheet(db,path,kind):
 from openpyxl import load_workbook
 wb=load_workbook(path,data_only=True,read_only=True);name='Materials' if kind=='materials' else 'Samples'
 if name not in wb.sheetnames:raise ValueError(f"Workbook has no {name} sheet")
 rows=wb[name].iter_rows(values_only=True);headers=[str(x) for x in next(rows)];added=skipped=0;warnings=[]
 for number,values in enumerate(rows,2):
  row=dict(zip(headers,values))
  try:
   if not row.get('name'):raise ValueError('name required')
   (db.save_material if kind=='materials' else db.save_sample)(row);added+=1
  except Exception as e:skipped+=1;warnings.append(f"Row {number}: {e}")
 return added,skipped,warnings
